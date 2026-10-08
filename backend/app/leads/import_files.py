"""Extract addresses from uploaded files without fixed file or row limits."""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import BinaryIO
from xml.etree import ElementTree as ET

from app.email.email_cleaner import is_acceptable_email

_EMAIL = re.compile(
    r"(?<![A-Za-z0-9._%+-])([A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,})"
    r"(?![A-Za-z0-9.-])", re.I,
)
_TEXT_FORMATS = {".csv", ".tsv", ".txt", ".html", ".htm", ".rtf"}


class ImportFileError(ValueError):
    pass


def _text_lines(source: BinaryIO) -> Iterator[str]:
    source.seek(0)
    marker = source.read(4)
    source.seek(0)
    encoding = "utf-16" if marker.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    wrapper = io.TextIOWrapper(source, encoding=encoding, errors="replace")
    try:
        yield from wrapper
    finally:
        wrapper.detach()


def _xml_paragraphs(source: BinaryIO, member: str) -> Iterator[str]:
    try:
        source.seek(0)
        with zipfile.ZipFile(source) as archive:
            if member not in archive.namelist():
                raise ImportFileError("Document has no readable text")
            with archive.open(member) as xml:
                for _, element in ET.iterparse(xml, events=("end",)):
                    if element.tag.endswith("}p"):
                        yield "".join(element.itertext())
                        element.clear()
    except zipfile.BadZipFile as exc:
        raise ImportFileError("Invalid spreadsheet or document file") from exc
    except ET.ParseError as exc:
        raise ImportFileError("Document text is malformed") from exc


def _file_texts(source: BinaryIO, suffix: str) -> Iterator[str]:
    if suffix in _TEXT_FORMATS or not suffix:
        yield from _text_lines(source)
        return
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook

            source.seek(0)
            book = load_workbook(source, read_only=True, data_only=True)
            try:
                for sheet in book.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        for value in row:
                            if value is not None:
                                yield str(value)
            finally:
                book.close()
        except Exception as exc:
            raise ImportFileError("Could not read XLSX workbook") from exc
        return
    if suffix == ".xls":
        try:
            import xlrd

            source.seek(0)
            book = xlrd.open_workbook(file_contents=source.read(), on_demand=True)
            try:
                for sheet in book.sheets():
                    for row in range(sheet.nrows):
                        for value in sheet.row_values(row):
                            if value:
                                yield str(value)
            finally:
                book.release_resources()
        except Exception as exc:
            raise ImportFileError("Could not read XLS workbook") from exc
        return
    if suffix == ".pdf":
        import pdfplumber

        try:
            source.seek(0)
            with pdfplumber.open(source) as pdf:
                for page in pdf.pages:
                    yield page.extract_text() or ""
        except Exception as exc:
            raise ImportFileError("Could not read PDF text; scanned PDFs need OCR") from exc
        return
    if suffix == ".docx":
        yield from _xml_paragraphs(source, "word/document.xml")
        return
    if suffix == ".ods":
        yield from _xml_paragraphs(source, "content.xml")
        return
    raise ImportFileError("Unsupported file; use CSV, XLSX, XLS, PDF, DOCX, ODS or text")


def extract_emails(data: bytes | BinaryIO, filename: str) -> tuple[list[str], int]:
    """Return unique acceptable emails and count of rejected/duplicate hits."""
    source = io.BytesIO(data) if isinstance(data, bytes) else data
    source.seek(0, io.SEEK_END)
    if source.tell() == 0:
        raise ImportFileError("File is empty")
    source.seek(0)
    seen: set[str] = set()
    emails: list[str] = []
    rejected = 0
    for text in _file_texts(source, Path(filename or "upload.txt").suffix.lower()):
        for match in _EMAIL.finditer(text):
            email = match.group(1).lower()
            if email in seen or not is_acceptable_email(email):
                rejected += 1
                continue
            seen.add(email)
            emails.append(email)
    if not emails:
        raise ImportFileError("No usable email addresses found in this file")
    return emails, rejected
